from fpdf import FPDF
import PyPDF2
import docx
import io
import os
import re

def extract_text_from_pdf(file):
    reader = PyPDF2.PdfReader(file)
    text = ""
    for page in reader.pages:
        text += page.extract_text() + "\n"
    return text.strip()

def extract_text_from_docx(file):
    doc = docx.Document(file)
    return "\n".join([para.text for para in doc.paragraphs])

def read_resume(file):
    if file.name.endswith(".pdf"):
        return extract_text_from_pdf(file)
    elif file.name.endswith(".docx"):
        return extract_text_from_docx(file)
    else:
        return None
    
def text_to_pdf_bytes(text: str, title: str = "Document") -> bytes:
    """
    Compact, block-aware renderer with keep-together rules and a hard 2-page cap.
    - Headings stick with first subheading
    - Subheadings stick with at least the first bullet
    - Measures height before drawing to avoid mid-sentence/page cutoffs
    - Trims gracefully if content would exceed 2 pages
    """

    # ---------------- Compact defaults (tweak if you like) ----------------
    MARGINS_MM = (16, 16, 16)   # left, top, right
    BODY_SIZE = 10.5
    HEAD_SIZE = 11.5
    SUBHEAD_SIZE = 11
    LINE_H_BODY = 5.5
    LINE_H_HEAD = 6.0
    GAP_AFTER_PARA = 1.0
    GAP_AFTER_SUBHEAD = 2.0
    GAP_AFTER_HEADING = 2.0
    BULLET_INDENT_MM = 3.0
    CONTACT_SIZE = 13
    MAX_PAGES = 2

    # ---------------------------------------------------------------
    text = text.replace("\t", "    ").replace("—", "–")  # normalize

    # Fonts
    font_reg = os.path.join("fonts", "DejaVuSans.ttf")
    font_bold = os.path.join("fonts", "DejaVuSans-Bold.ttf")
    font_oblq = os.path.join("fonts", "DejaVuSans-Oblique.ttf")

    pdf = FPDF()
    l, t, r = MARGINS_MM
    pdf.set_margins(l, t, r)
    pdf.add_page()

    pdf.add_font("DejaVu", "", font_reg, uni=True)
    pdf.add_font("DejaVu", "B", font_bold, uni=True)
    pdf.add_font("DejaVuI", "", font_oblq, uni=True)

    pdf.set_font("DejaVu", "", BODY_SIZE)

    page_w = pdf.w - pdf.l_margin - pdf.r_margin
    bullet_re = re.compile(r"^(\s*)([●•\-\*])\s+(.*)$")

    SECTION_HINTS = {
        "SUMMARY", "PROFILE", "KEY SKILLS", "SKILLS", "CAPABILITIES",
        "EDUCATION", "PROJECTS", "SELECTED PROJECTS",
        "EXPERIENCE", "WORK HISTORY", "ADDITIONAL", "INTERESTS"
    }

    def is_section_heading(s: str) -> bool:
        s = s.strip()
        if not s:
            return False
        if s.upper() in SECTION_HINTS:
            return True
        if s.isupper() and 3 <= len(s) <= 40:
            return True
        if s.endswith(":") and len(s) <= 40:
            return True
        return False

    def is_subheading(s: str) -> bool:
        s = s.strip()
        if not s or s.isupper() or bullet_re.match(s) or is_section_heading(s):
            return False
        if " | " in s or "," in s:
            return True
        if re.search(r"\b(19|20)\d{2}\b", s):
            return True
        if re.search(r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)\b", s, re.I):
            return True
        words = s.split()
        return 5 <= len(words) <= 16 and sum(w[:1].isupper() for w in words)/len(words) > 0.5

    def string_w(s: str) -> float:
        return pdf.get_string_width(s)

    def mm_to_user(u_mm: float) -> float:
        # FPDF units are mm by default; if changed, adjust accordingly.
        return u_mm

    bullet_indent = mm_to_user(BULLET_INDENT_MM)

    # --------- Word wrap estimator (mirrors MultiCell roughly) ----------
    def wrap_lines(text_line: str, avail_w: float) -> int:
        """Return estimated wrapped line count for a normal line."""
        words = text_line.split(" ")
        if not words:
            return 1
        lines = 1
        curr = 0.0
        space_w = string_w(" ")
        for w in words:
            w_w = string_w(w)
            if curr == 0:
                curr = w_w
            elif curr + space_w + w_w <= avail_w:
                curr += space_w + w_w
            else:
                lines += 1
                curr = w_w
        return max(lines, 1)

    def bullet_lines(rest_text: str, avail_first: float, avail_other: float) -> int:
        """Wrapped lines for a hanging bullet: first line width differs."""
        words = rest_text.split(" ")
        if not words:
            return 1
        lines = 1
        curr = 0.0
        space_w = string_w(" ")
        width = avail_first
        first_line_done = False
        for w in words:
            w_w = string_w(w)
            if curr == 0:
                curr = w_w
            elif curr + space_w + w_w <= width:
                curr += space_w + w_w
            else:
                lines += 1
                curr = w_w
                if not first_line_done:
                    first_line_done = True
                    width = avail_other
        return max(lines, 1)

    # ----------------- Parse into blocks (keep-together units) -----------------
    lines = text.splitlines()

    # Identify first non-empty as contact line
    first_non_empty_idx = next((i for i, ln in enumerate(lines) if ln.strip()), None)

    blocks = []  # each block = dict(type, lines, meta)
    i = 0
    while i < len(lines):
        ln = lines[i]

        # Blank line -> spacing block
        if ln.strip() == "":
            blocks.append({"type": "spacer", "lines": [ln]})
            i += 1
            continue

        # Contact (first non-empty)
        if i == first_non_empty_idx:
            blocks.append({"type": "contact", "lines": [ln.strip()]})
            i += 1
            continue

        # Section heading
        if is_section_heading(ln):
            # heading + (optional next subheading kept with it)
            block = {"type": "heading", "lines": [ln.strip()]}
            # Peek next non-empty for subheading; we don't *merge* it here,
            # but we will enforce keep-with-next when placing.
            blocks.append(block)
            i += 1
            continue

        # Subheading and its following bullet/paragraph group
        if is_subheading(ln):
            sub = [ln.strip()]
            i += 1
            # Collect following bullets/paragraphs until blank line or next heading/subheading
            group = []
            while i < len(lines):
                nxt = lines[i]
                if nxt.strip() == "":
                    break
                if is_section_heading(nxt) or is_subheading(nxt):
                    break
                group.append(nxt.rstrip())
                i += 1
            blocks.append({"type": "subheading", "lines": sub})
            if group:
                # group may contain bullets and normal lines; keep together
                blocks.append({"type": "group", "lines": group})
            continue

        # Any other single paragraph line (kept as its own block)
        blocks.append({"type": "para", "lines": [ln.rstrip()]})
        i += 1

    # -------------- Measurement helpers (block heights before drawing) --------------
    def height_of_block(block) -> float:
        btype = block["type"]
        if btype == "spacer":
            return LINE_H_BODY  # one blank line
        if btype == "contact":
            # one centered line + a small gap
            return (LINE_H_BODY + 1) + LINE_H_BODY
        if btype == "heading":
            pdf.set_font("DejaVu", "B", HEAD_SIZE)
            h = 0.0
            for s in block["lines"]:
                h += LINE_H_HEAD * wrap_lines(s.strip(), page_w)
            h += GAP_AFTER_HEADING
            pdf.set_font("DejaVu", "", BODY_SIZE)
            return h
        if btype == "subheading":
            pdf.set_font("DejaVuI", "", SUBHEAD_SIZE)
            h = 0.0
            for s in block["lines"]:
                h += LINE_H_HEAD * wrap_lines(s.strip(), page_w)
            h += GAP_AFTER_SUBHEAD
            pdf.set_font("DejaVu", "", BODY_SIZE)
            return h
        if btype in ("para", "group"):
            h = 0.0
            for s in block["lines"]:
                m = bullet_re.match(s)
                if m:
                    leading, sym, rest = m.groups()
                    indent_w = string_w(leading)
                    bullet_token = f"{sym} "
                    token_w = string_w(bullet_token)
                    first_avail = page_w - indent_w - token_w - string_w(" ")
                    other_avail = page_w - indent_w - token_w - string_w(" ")
                    lines_cnt = bullet_lines(rest, first_avail, other_avail)
                    h += lines_cnt * LINE_H_BODY + GAP_AFTER_PARA
                else:
                    leading_spaces = len(s) - len(s.lstrip(" "))
                    indent_w = string_w(" " * leading_spaces)
                    avail = page_w - indent_w
                    lines_cnt = wrap_lines(s.strip(), avail)
                    h += lines_cnt * LINE_H_BODY + GAP_AFTER_PARA
            return h
        return LINE_H_BODY

    # ----------------- Placement with keep-together & 2-page cap --------------------
    def new_page():
        pdf.add_page()
        pdf.set_font("DejaVu", "", BODY_SIZE)

    # Keep-with-next rules will be enforced on placement:
    # - heading keeps with next subheading if present
    # - subheading keeps with at least first line of next group

    page_count = 1
    idx = 0
    appended_footer = False

    def remaining_height() -> float:
        return pdf.h - pdf.b_margin - pdf.get_y()

    while idx < len(blocks):
        blk = blocks[idx]
        btype = blk["type"]

        # Build a composite unit respecting keep-with-next
        unit = [blk]
        # If heading, peek next subheading to keep together
        if btype == "heading" and idx + 1 < len(blocks) and blocks[idx + 1]["type"] == "subheading":
            unit.append(blocks[idx + 1])
            # If also followed by a group, keep at least first line of it
            if idx + 2 < len(blocks) and blocks[idx + 2]["type"] == "group":
                # Split group into first line block + remainder
                grp_lines = blocks[idx + 2]["lines"]
                if grp_lines:
                    unit.append({"type": "group", "lines": [grp_lines[0]]})
        # If subheading, keep with at least first line of next group
        elif btype == "subheading" and idx + 1 < len(blocks) and blocks[idx + 1]["type"] == "group":
            first_only = {"type": "group", "lines": [blocks[idx + 1]["lines"][0]]}
            unit.append(first_only)

        # Measure composite height
        unit_h = sum(height_of_block(b) for b in unit)

        # If it doesn't fit on current page, go to next page
        if unit_h > remaining_height():
            # If already on page 2, we need to trim (enforce cap)
            if page_count >= MAX_PAGES:
                # Try a soft footer once
                if not appended_footer:
                    pdf.ln(LINE_H_BODY)
                    pdf.set_font("DejaVu", "B", BODY_SIZE)
                    pdf.multi_cell(page_w, LINE_H_BODY, "(Additional details available on request.)", align="L")
                    pdf.set_font("DejaVu", "", BODY_SIZE)
                    appended_footer = True
                break
            # New page
            new_page()
            page_count += 1
            continue

        # It fits: draw the actual block (not the reduced unit preview)
        # Draw current block
        def draw_block(block):
            btype = block["type"]
            if btype == "spacer":
                pdf.ln(LINE_H_BODY)
                return
            if btype == "contact":
                pdf.set_font("DejaVu", "B", CONTACT_SIZE)
                pdf.multi_cell(page_w, LINE_H_BODY + 1, block["lines"][0], align="C")
                pdf.ln(LINE_H_BODY)
                pdf.set_font("DejaVu", "", BODY_SIZE)
                return
            if btype == "heading":
                pdf.set_font("DejaVu", "B", HEAD_SIZE)
                for s in block["lines"]:
                    pdf.multi_cell(page_w, LINE_H_HEAD, s.strip(), align="L")
                pdf.ln(GAP_AFTER_HEADING)
                pdf.set_font("DejaVu", "", BODY_SIZE)
                return
            if btype == "subheading":
                pdf.set_font("DejaVuI", "", SUBHEAD_SIZE)
                for s in block["lines"]:
                    pdf.multi_cell(page_w, LINE_H_HEAD, s.strip(), align="L")
                pdf.ln(GAP_AFTER_SUBHEAD)
                pdf.set_font("DejaVu", "", BODY_SIZE)
                return
            if btype in ("para", "group"):
                for s in block["lines"]:
                    m = bullet_re.match(s)
                    if m:
                        leading, sym, rest = m.groups()
                        indent_w = pdf.get_string_width(leading)
                        y0 = pdf.get_y()
                        x_bullet = pdf.l_margin + indent_w
                        pdf.set_xy(x_bullet, y0)
                        # smaller bullet
                        pdf.set_font("DejaVu", "", BODY_SIZE - 1)
                        token = f"{sym} "
                        pdf.cell(pdf.get_string_width(token), LINE_H_BODY, token, ln=0)
                        # hanging indent
                        pdf.set_font("DejaVu", "", BODY_SIZE)
                        x_text = pdf.get_x() + pdf.get_string_width(" ")
                        pdf.set_xy(x_text, y0)
                        avail = pdf.w - pdf.r_margin - x_text
                        pdf.multi_cell(avail, LINE_H_BODY, rest, align="J")
                        pdf.ln(GAP_AFTER_PARA)
                    else:
                        leading_spaces = len(s) - len(s.lstrip(" "))
                        indent_w = pdf.get_string_width(" " * leading_spaces)
                        pdf.set_x(pdf.l_margin + indent_w)
                        pdf.multi_cell(page_w - indent_w, LINE_H_BODY, s.strip(), align="J")
                        pdf.ln(GAP_AFTER_PARA)
                return

        # Draw the full block; also draw next kept pieces if we previewed them
        draw_block(blk)

        # If heading and next is subheading, ensure we draw it now
        advanced = 1
        if btype == "heading" and idx + 1 < len(blocks) and blocks[idx + 1]["type"] == "subheading":
            draw_block(blocks[idx + 1])
            advanced += 1
            if idx + 2 < len(blocks) and blocks[idx + 2]["type"] == "group":
                # draw just the first line of the group here
                first_line = blocks[idx + 2]["lines"][0]
                draw_block({"type": "group", "lines": [first_line]})
                # mutate the original group to drop the first line (so it will render next)
                blocks[idx + 2]["lines"] = blocks[idx + 2]["lines"][1:] or []
                if not blocks[idx + 2]["lines"]:
                    advanced += 1  # consumed entire group

        elif btype == "subheading" and idx + 1 < len(blocks) and blocks[idx + 1]["type"] == "group":
            # draw first line of the following group now
            first_line = blocks[idx + 1]["lines"][0]
            draw_block({"type": "group", "lines": [first_line]})
            # mutate that group to drop the first line
            blocks[idx + 1]["lines"] = blocks[idx + 1]["lines"][1:] or []
            if not blocks[idx + 1]["lines"]:
                advanced += 1

        idx += advanced

        # If we just filled second page to the brim, add footer if anything remains
        if page_count >= MAX_PAGES and idx < len(blocks) and remaining_height() <= LINE_H_BODY * 2 and not appended_footer:
            pdf.ln(LINE_H_BODY)
            pdf.set_font("DejaVu", "B", BODY_SIZE)
            pdf.multi_cell(page_w, LINE_H_BODY, "(Additional details available on request.)", align="L")
            pdf.set_font("DejaVu", "", BODY_SIZE)
            appended_footer = True
            break

    return pdf.output(dest="S").encode("latin1")